from pathlib import Path
from docx import Document
from docx.shared import Inches
import matplotlib.pyplot as plt


class BusinessReport:

    @staticmethod
    def generate_reports(
            execution_results,
            total_execution_time):

        Path(
            "reports/business"
        ).mkdir(
            parents=True,
            exist_ok=True
        )

        BusinessReport.generate_summary_report(
            execution_results,
            total_execution_time
        )

        BusinessReport.generate_testcase_reports(
            execution_results
        )

    @staticmethod
    def create_pie_chart(
            passed,
            failed):

        plt.figure(
            figsize=(5, 5)
        )

        plt.pie(
            [passed, failed],
            labels=["PASS", "FAIL"],
            autopct="%1.1f%%",
            colors=["green", "red"]
        )

        plt.title(
            "Execution Summary"
        )

        chart_path = (
            "reports/business/"
            "execution_summary_chart.png"
        )

        plt.savefig(
            chart_path,
            bbox_inches="tight"
        )

        plt.close()

        return chart_path

    @staticmethod
    def generate_summary_report(
            execution_results,
            total_execution_time):

        document = Document()

        document.add_heading(
            "TEST EXECUTION DASHBOARD",
            level=1
        )

        total = len(
            execution_results
        )

        passed = len(
            [
                tc
                for tc in execution_results
                if tc["status"] == "PASS"
            ]
        )

        failed = len(
            [
                tc
                for tc in execution_results
                if tc["status"] == "FAIL"
            ]
        )

        pass_percentage = (
            round(
                (passed / total) * 100,
                2
            )
            if total > 0
            else 0
        )

        fail_percentage = (
            round(
                (failed / total) * 100,
                2
            )
            if total > 0
            else 0
        )

        overall_status = (
            "GO"
            if failed == 0
            else "NO GO"
        )

        document.add_heading(
            "Execution Summary",
            level=2
        )

        document.add_paragraph(
            f"Total Test Cases : {total}"
        )

        document.add_paragraph(
            f"Passed : {passed}"
        )

        document.add_paragraph(
            f"Failed : {failed}"
        )

        document.add_paragraph(
            f"Pass Percentage : "
            f"{pass_percentage}%"
        )

        document.add_paragraph(
            f"Fail Percentage : "
            f"{fail_percentage}%"
        )

        document.add_paragraph(
            f"Overall Execution Time : "
            f"{total_execution_time} sec"
        )

        document.add_paragraph(
            f"Overall Status : "
            f"{overall_status}"
        )

        document.add_paragraph("")

        chart_path = (
            BusinessReport.create_pie_chart(
                passed,
                failed
            )
        )

        document.add_picture(
            chart_path,
            width=Inches(3)
        )

        document.add_page_break()

        document.add_heading(
            "Execution Results",
            level=2
        )

        table = document.add_table(
            rows=1,
            cols=4
        )

        table.style = (
            "Table Grid"
        )

        header = table.rows[0].cells

        header[0].text = (
            "Test Case ID"
        )
        header[1].text = (
            "Test Case Name"
        )
        header[2].text = (
            "Status"
        )
        header[3].text = (
            "Duration (Sec)"
        )

        for tc in execution_results:

            row = (
                table.add_row().cells
            )

            row[0].text = (
                tc["testcase_id"]
            )

            row[1].text = (
                tc["testcase_name"]
            )

            row[2].text = (
                tc["status"]
            )

            row[3].text = str(
                tc.get(
                    "duration",
                    ""
                )
            )

        document.add_page_break()

        document.add_heading(
            "Failed Scenarios",
            level=2
        )

        failed_found = False

        for tc in execution_results:

            if tc["status"] == "FAIL":

                failed_found = True

                document.add_paragraph(
                    f"Test Case ID : "
                    f"{tc['testcase_id']}"
                )

                document.add_paragraph(
                    f"Test Case Name : "
                    f"{tc['testcase_name']}"
                )

                document.add_paragraph(
                    f"Reason : "
                    f"{tc['failure_reason']}"
                )

                document.add_paragraph("")

        if not failed_found:

            document.add_paragraph(
                "No failed scenarios."
            )

        document.add_page_break()

        document.add_heading(
            "Recommendation",
            level=2
        )

        document.add_paragraph(
            overall_status
        )

        document.save(
            "reports/business/"
            "Execution_Summary.docx"
        )

    @staticmethod
    def generate_testcase_reports(
            execution_results):

        for tc in execution_results:

            document = Document()

            document.add_heading(
                "TEST CASE EVIDENCE REPORT",
                level=1
            )

            document.add_paragraph(
                f"Test Case ID : "
                f"{tc['testcase_id']}"
            )

            document.add_paragraph(
                f"Test Case Name : "
                f"{tc['testcase_name']}"
            )

            document.add_paragraph(
                f"Description : "
                f"{tc['description']}"
            )

            document.add_paragraph(
                f"Environment : "
                f"{tc['environment']}"
            )

            document.add_paragraph(
                f"Status : "
                f"{tc['status']}"
            )

            document.add_paragraph(
                f"Execution Time : "
                f"{tc.get('duration', '')}"
                f" sec"
            )

            document.add_paragraph("")

            for index, step in enumerate(
                    tc["steps"],
                    start=1):

                document.add_heading(
                    f"Step {index}",
                    level=2
                )

                document.add_paragraph(
                    f"Step Name : "
                    f"{step['step_name']}"
                )

                document.add_paragraph(
                    f"Description : "
                    f"{step['description']}"
                )

                document.add_paragraph(
                    f"Validation : "
                    f"{step['validation']}"
                )

                screenshot = (
                    step["screenshot"]
                )

                if screenshot:

                    screenshot_path = Path(
                        screenshot
                    )

                    if screenshot_path.exists():

                        document.add_picture(
                            str(
                                screenshot_path
                            ),
                            width=Inches(5)
                        )

                document.add_paragraph("")

            if tc["status"] == "FAIL":

                document.add_heading(
                    "Failure Details",
                    level=2
                )

                document.add_paragraph(
                    tc["failure_reason"]
                )

            safe_name = (
                tc["testcase_name"]
                .replace(" ", "_")
                .replace("/", "_")
            )

            filename = (
                f"{tc['testcase_id']}_"
                f"{safe_name}.docx"
            )

            document.save(
                f"reports/business/"
                f"{filename}"
            )